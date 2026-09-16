#!/usr/bin/env python3
from __future__ import annotations

import re
import subprocess
from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "research" / "t6" / "build"
GEN = OUT / "generated"
OUT.mkdir(parents=True, exist_ok=True)
GEN.mkdir(parents=True, exist_ok=True)

PARTS = [
    ROOT / "docs/t6/T6_INTRODUCTION.md",
    ROOT / "docs/t6/T6_CHAPTER_1.md",
    ROOT / "docs/t6/T6_CHAPTER_2.md",
    ROOT / "docs/t6/T6_CHAPTER_3.md",
    ROOT / "docs/t6/T6_CONCLUSION.md",
    ROOT / "docs/t6/T6_REFERENCES.md",
]


def dot(name: str, body: str) -> Path:
    src = GEN / f"{name}.dot"
    png = GEN / f"{name}.png"
    src.write_text(
        'digraph G {\nrankdir=LR; graph [bgcolor=white,pad=0.2,nodesep=0.45,ranksep=0.6]; '
        'node [shape=box,style="rounded",fontname="DejaVu Sans",fontsize=10,margin=0.12]; '
        'edge [fontname="DejaVu Sans",fontsize=9];\n' + body + '\n}\n',
        encoding="utf-8",
    )
    subprocess.run(["dot", "-Tpng", "-Gdpi=160", str(src), "-o", str(png)], check=True)
    return png


def make_diagrams() -> list[Path]:
    return [
        dot("tobe", '''A[label="Получение"]; B[label="Регистрация + provenance"]; C[label="Валидация / dedup / обезличивание"]; D[label="Классификация\\nтональность + аспекты + C1-C4"]; E[label="C3/C4?",shape=diamond]; F[label="Expert verification"]; G[label="Стандартная маршрутизация"]; H[label="Эскалация"]; I[label="Профильная роль"]; J[label="Решение"]; K[label="Ответ клиенту"]; L[label="Внутреннее мероприятие"]; M[label="Контроль"]; N[label="Outcome"]; O[label="Результат подтверждён?",shape=diamond]; P[label="Закрытие + знания"]; Q[label="Аналитика / KPI"]; A->B->C->D->E; E->F[label=" да"]; E->G[label=" нет"]; F->H->J; G->I->J; J->K->M; J->L->M; M->N->O; O->J[label=" нет"]; O->P[label=" да"]; P->Q;'''),
        dot("decision_tree", '''A[label="Новый отзыв"]; B[label="Регистрация"]; C[label="Аспекты + тональность"]; D[label="Critical trigger?",shape=diamond]; E[label="Предварительный C3/C4"]; F[label="Expert verification"]; G[label="Высокий риск подтверждён?",shape=diamond]; H[label="Эскалация RM/LC"]; I[label="Изменить C-level\\nс основанием"]; J[label="Нужно профильное действие?",shape=diamond]; K[label="C2/C3 по контексту"]; L[label="C1"]; M[label="Решение"]; N[label="Ответ и/или мероприятие"]; O[label="Контроль"]; P[label="Outcome подтверждён?",shape=diamond]; Q[label="Закрытие"]; A->B->C->D; D->E[label=" да"]; E->F->G; G->H[label=" да"]; G->I[label=" нет"]; D->J[label=" нет"]; J->K[label=" да"]; J->L[label=" нет"]; H->M; I->M; K->M; L->M; M->N->O->P; P->M[label=" нет"]; P->Q[label=" да"];'''),
        dot("architecture", '''A[label="Каналы обратной связи"]; B[label="Получение и регистрация"]; C[label="Качество данных"]; D[label="Классификация и анализ"]; E[label="Решение и маршрутизация"]; F[label="Коммуникация"]; G[label="Мероприятия"]; H[label="Контроль outcome"]; I[label="Аналитика / KPI"]; J[label="Knowledge base / улучшения"]; A->B->C->D->E; E->F->H; E->G->H; H->I->J->D;'''),
        dot("logical_model", '''S[label="SOURCE"]; R[label="REVIEW_CASE"]; C[label="CLASSIFICATION"]; A[label="ASPECT"]; D[label="DECISION"]; RESP[label="RESPONSE"]; ACT[label="ACTION"]; CTRL[label="CONTROL_EVENT"]; O[label="OUTCOME"]; ROLE[label="ROLE"]; AUD[label="AUDIT_EVENT"]; S->R; R->C; C->A[label="M:N"]; R->D; D->RESP; D->ACT; ACT->CTRL; RESP->CTRL; R->O; ROLE->D; ROLE->ACT; R->AUD;'''),
        dot("implementation", '''E0[label="E0 Решение о пилоте"]; E1[label="E1 Подготовка"]; E2[label="E2 Адаптация"]; E3[label="E3 Обучение / dry-run"]; E4[label="E4 BEFORE"]; E5[label="E5 Пилот"]; E6[label="E6 AFTER"]; E7[label="E7 Оценка"]; D[label="Решение",shape=diamond]; E8[label="E8 Масштабирование"]; X[label="Завершение / разбор причин"]; E0->E1->E2->E3->E4->E5->E6->E7->D; D->E8[label=" масштабировать"]; D->E2[label=" корректировать"]; D->X[label=" остановить"];'''),
    ]


def build_markdown(diagrams: list[Path]) -> Path:
    text = '''# ВЫПУСКНАЯ КВАЛИФИКАЦИОННАЯ РАБОТА\n\n**Тема:** «создание программы по работе с отзывами клиентов анализ, моделирование, внедрение»\n\n**Рабочее название авторской разработки:** «Методика ЕГ»\n\n*Рабочая сборка T6. Реквизиты образовательной организации, обучающегося и руководителя заполняются по утверждённому титульному шаблону на этапе T7.*\n\n'''
    captions = [
        "Целевая процессная модель TO-BE «Методики ЕГ»",
        "Дерево решений и эскалации",
        "Функциональная архитектура «Методики ЕГ»",
        "Логическая модель информационных сущностей",
        "Этапы внедрения E0-E8",
    ]
    d_idx = 0
    for path in PARTS:
        s = path.read_text(encoding="utf-8-sig").replace("../../research/", "research/")
        def repl(_m):
            nonlocal d_idx
            if d_idx >= len(diagrams):
                return ""
            img = diagrams[d_idx].relative_to(ROOT).as_posix()
            cap = captions[d_idx]
            d_idx += 1
            return f"\n![{cap}]({img})\n"
        s = re.sub(r"```mermaid\s*.*?```", repl, s, flags=re.S)
        text += s.rstrip() + "\n\n"
    combined = OUT / "T6_COMBINED.md"
    combined.write_text(text, encoding="utf-8")
    return combined


def set_run_font(run, size=Pt(14)):
    run.font.name = "Times New Roman"
    run.font.size = size
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{attr}"), "Times New Roman")


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar"); begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve"); instr.text = " PAGE "
    end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, end])
    set_run_font(run, Pt(12))


def add_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = borders.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "auto")


def postprocess(docx_path: Path):
    doc = Document(docx_path)
    for sec in doc.sections:
        sec.top_margin, sec.bottom_margin = Cm(2.0), Cm(2.0)
        sec.left_margin, sec.right_margin = Cm(3.0), Cm(1.5)
        sec.header_distance, sec.footer_distance = Cm(1.0), Cm(1.0)
        fp = sec.footer.paragraphs[0] if sec.footer.paragraphs else sec.footer.add_paragraph()
        if not fp.text:
            add_page_number(fp)

    for stylename in ["Normal", "Body Text", "First Paragraph"]:
        if stylename in doc.styles:
            st = doc.styles[stylename]
            st.font.name, st.font.size = "Times New Roman", Pt(14)
            st.paragraph_format.line_spacing = 1.5
            st.paragraph_format.space_after = Pt(0)
            st.paragraph_format.first_line_indent = Cm(1.25)
    for stylename, size in [("Title", 16), ("Heading 1", 14), ("Heading 2", 14), ("Heading 3", 14)]:
        if stylename in doc.styles:
            st = doc.styles[stylename]
            st.font.name, st.font.size, st.font.bold = "Times New Roman", Pt(size), True
            st.paragraph_format.first_line_indent = Cm(0)
            st.paragraph_format.space_before, st.paragraph_format.space_after = Pt(12), Pt(6)
            st.paragraph_format.keep_with_next = True

    for p in doc.paragraphs:
        txt = p.text.strip()
        style = p.style.name if p.style else ""
        heading = style.startswith("Heading") or style == "Title"
        caption = txt.startswith(("Рисунок ", "Таблица ", "Источник:"))
        if heading:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if style in ("Heading 1", "Title") else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.first_line_indent = Cm(0)
            if style == "Heading 1" and txt != "ВЫПУСКНАЯ КВАЛИФИКАЦИОННАЯ РАБОТА":
                p.paragraph_format.page_break_before = True
        elif txt:
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.line_spacing = 1.5
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.first_line_indent = Cm(0) if caption else Cm(1.25)
        if p._p.xpath('.//w:drawing'):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.first_line_indent = Cm(0)
        for r in p.runs:
            set_run_font(r, Pt(12) if caption else Pt(14))

    for table in doc.tables:
        add_table_borders(table)
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    p.paragraph_format.first_line_indent = Cm(0)
                    p.paragraph_format.line_spacing = 1.0
                    p.paragraph_format.space_after = Pt(0)
                    for r in p.runs:
                        set_run_font(r, Pt(10))
        if table.rows:
            tr_pr = table.rows[0]._tr.get_or_add_trPr()
            header = OxmlElement("w:tblHeader"); header.set(qn("w:val"), "true"); tr_pr.append(header)

    doc.core_properties.title = "ВКР — Методика ЕГ — T6"
    doc.core_properties.subject = "Создание программы по работе с отзывами клиентов анализ, моделирование, внедрение"
    doc.core_properties.author = ""
    doc.save(docx_path)


def main():
    md = build_markdown(make_diagrams())
    final = OUT / "WKR_Methodika_EG_T6.docx"
    subprocess.run([
        "pandoc", str(md), "--from=markdown", "--to=docx", "--toc", "--toc-depth=2",
        "--resource-path", str(ROOT), "-o", str(final)
    ], cwd=ROOT, check=True)
    postprocess(final)
    print(final)

if __name__ == "__main__":
    main()
