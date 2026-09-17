#!/usr/bin/env python3
from __future__ import annotations

import csv
import json
import subprocess
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "research/t7/build_appendices"
GEN = OUT / "generated"
OUT.mkdir(parents=True, exist_ok=True)
GEN.mkdir(parents=True, exist_ok=True)
DOCX = OUT / "WKR_Methodika_EG_Appendices_A_T.docx"

TABLE_DIR = ROOT / "research/t3/output/tables"
CHART_DIR = ROOT / "research/t3/output/charts"
MANIFEST = ROOT / "research/t3/output/acquisition_manifest.json"
P1P8 = ROOT / "research/t3/output/P1_P8_MATRIX.csv"

LETTERS = ["А", "Б", "В", "Г", "Д", "Е", "Ж", "И", "К", "Л", "М", "Н", "П", "Р", "С", "Т"]


def set_cell_text(cell, text: str, size=9, bold=False):
    cell.text = str(text)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for p in cell.paragraphs:
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.space_before = Pt(0)
        for r in p.runs:
            r.font.name = "Times New Roman"
            r._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
            r.font.size = Pt(size)
            r.bold = bold


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_table(doc: Document, rows: list[list[str]], header=True, font_size=9, widths_cm=None):
    if not rows:
        return None
    cols = max(len(r) for r in rows)
    table = doc.add_table(rows=0, cols=cols)
    table.style = "Table Grid"
    table.autofit = False if widths_cm else True
    for ri, row in enumerate(rows):
        cells = table.add_row().cells
        for ci in range(cols):
            value = row[ci] if ci < len(row) else ""
            set_cell_text(cells[ci], value, font_size, bold=header and ri == 0)
            if widths_cm and ci < len(widths_cm):
                cells[ci].width = Cm(widths_cm[ci])
        if header and ri == 0:
            set_repeat_table_header(table.rows[0])
    doc.add_paragraph()
    return table


def add_body(doc: Document, text: str, bold=False, italic=False):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(1.25)
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(0)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    r.font.size = Pt(14)
    r.bold = bold
    r.italic = italic
    return p


def add_bullets(doc: Document, items: list[str]):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Cm(1.0)
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(item)
        r.font.name = "Times New Roman"
        r._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        r.font.size = Pt(12)


def section_portrait(section):
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Mm(20)
    section.bottom_margin = Mm(20)
    section.left_margin = Mm(30)
    section.right_margin = Mm(15)


def section_landscape(section):
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Mm(297)
    section.page_height = Mm(210)
    section.top_margin = Mm(15)
    section.bottom_margin = Mm(15)
    section.left_margin = Mm(15)
    section.right_margin = Mm(15)


def new_section(doc: Document, landscape=False):
    section = doc.add_section(WD_SECTION.NEW_PAGE)
    if landscape:
        section_landscape(section)
    else:
        section_portrait(section)
    return section


def appendix_title(doc: Document, letter: str, title: str):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(12)
    r = p.add_run(f"ПРИЛОЖЕНИЕ {letter}\n{title}")
    r.font.name = "Times New Roman"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    r.font.size = Pt(14)
    r.bold = True


def add_caption(doc: Document, text: str):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    r.font.size = Pt(11)


def add_image(doc: Document, path: Path, width_cm=15.0):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run()
    r.add_picture(str(path), width=Cm(width_cm))


def dot_png(name: str, dot_body: str) -> Path:
    dot_path = GEN / f"{name}.dot"
    png_path = GEN / f"{name}.png"
    dot_path.write_text(
        "digraph G { graph [bgcolor=white, pad=0.2, nodesep=0.35, ranksep=0.5]; "
        "node [shape=box, style=\"rounded\", fontname=\"DejaVu Sans\", fontsize=10]; "
        "edge [fontname=\"DejaVu Sans\", fontsize=9]; " + dot_body + " }",
        encoding="utf-8",
    )
    subprocess.run(["dot", "-Tpng", "-Gdpi=180", str(dot_path), "-o", str(png_path)], check=True)
    return png_path


def csv_rows(path: Path) -> list[list[str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as f:
        return [[str(x) for x in row] for row in csv.reader(f)]


def build() -> None:
    doc = Document()
    section_portrait(doc.sections[0])
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(14)
    normal.paragraph_format.line_spacing = 1.5

    # Cover for the separate appendix packet.
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Cm(7)
    r = p.add_run("ПРИЛОЖЕНИЯ К ВЫПУСКНОЙ КВАЛИФИКАЦИОННОЙ РАБОТЕ")
    r.font.name = "Times New Roman"; r._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman"); r.font.size = Pt(16); r.bold = True
    p2 = doc.add_paragraph(); p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = p2.add_run("Авторская организационно-методическая разработка «Методика ЕГ»")
    r2.font.name = "Times New Roman"; r2._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman"); r2.font.size = Pt(14); r2.bold = True
    add_body(doc, "Пакет содержит только материалы, прослеживаемые к закрытой доказательной базе T3–T5. Приложения не подтверждают факт внедрения или проведения пилота.", italic=True)

    # A — corpus formation.
    new_section(doc)
    appendix_title(doc, "А", "Схема формирования эмпирического корпуса")
    manifest = json.loads(MANIFEST.read_text(encoding="utf-8"))
    assert manifest["selected_count"] == 1200
    assert manifest["selected_by_source"]["Yandex Maps"] == 1192
    assert manifest["selected_by_source"]["RuStore"] == 8
    add_body(doc, "Эмпирический корпус сформирован из публично доступных отзывов торговой сети «Перекрёсток» за период 01.09.2025–31.08.2026. Синтетическое дополнение данных не применялось.")
    add_table(doc, [
        ["Показатель", "Фактическое значение"],
        ["Исследовательский период", "01.09.2025–31.08.2026"],
        ["Пригодный пул до отбора", str(manifest["raw_eligible_before_dedup"])],
        ["Удалено дублей", str(manifest["duplicates_removed"])],
        ["Итоговый корпус", str(manifest["selected_count"])],
        ["Yandex Maps", str(manifest["selected_by_source"]["Yandex Maps"])],
        ["RuStore", str(manifest["selected_by_source"]["RuStore"])],
        ["Записи с рейтингом", str(manifest["selected_with_rating"])],
        ["Записи с видимым публичным ответом", str(manifest["selected_with_owner_reply"])],
    ], font_size=11, widths_cm=[8.5, 7.0])
    flow = dot_png("appendix_a_corpus", '''rankdir=LR;
A[label="Публичные источники\nYandex Maps + RuStore"] -> B[label="Проверка периода\nи пригодности"] -> C[label="Удаление дублей\nи обезличивание"] -> D[label="Пригодный пул\n1302 записи"] -> E[label="Детерминированный отбор"] -> F[label="Итоговый корпус\n1200 отзывов\n1192 + 8"];
''')
    add_image(doc, flow, 15.5); add_caption(doc, "Рисунок А.1 — Схема формирования корпуса")
    add_body(doc, "Открытые источники не позволяют наблюдать внутренние CRM, роли, SLA и корректирующие мероприятия компании; поэтому такие параметры не восстанавливаются предположениями.")

    # B — codebook.
    new_section(doc)
    appendix_title(doc, "Б", "Кодировочная схема анализа отзывов")
    add_body(doc, "Кодирование выполняется по независимым измерениям: прокси-показатель тональности, аспекты и критичность. Критичность не выводится автоматически из рейтинга.")
    add_table(doc, [
        ["Прокси тональности", "Правило"],
        ["POS", "4–5 звёзд"], ["NEU", "3 звезды"], ["NEG", "1–2 звезды"], ["UNKNOWN", "рейтинг отсутствует или ненадёжен"],
    ], font_size=11, widths_cm=[5, 10.5])
    aspects = [
        ("SERVICE", "обслуживание/сервис"), ("QUALITY", "качество товара/услуги"), ("STAFF", "персонал"),
        ("PRICE", "цена/акции/условия"), ("AVAILABILITY", "наличие/ассортимент"), ("SUPPORT", "поддержка/обратная связь"),
        ("DELIVERY", "доставка/сроки"), ("RETURN", "возврат/возмещение/претензия"), ("DIGITAL", "сайт/приложение/цифровой канал"), ("OTHER", "иная тема"),
    ]
    add_table(doc, [["Код", "Аспект"]] + [list(x) for x in aspects], font_size=10, widths_cm=[4, 11.5])
    add_table(doc, [
        ["Уровень", "Смысл"],
        ["C1", "локальное замечание без существенного ущерба"],
        ["C2", "требуется профильная функция или приоритетная обработка"],
        ["C3", "потенциально существенный ущерб, повторная нерешённая проблема, претензионный конфликт или высокий риск эскалации"],
        ["C4", "потенциальный критический правовой, безопасностный, массовый или репутационный эффект"],
    ], font_size=10, widths_cm=[3, 12.5])
    add_body(doc, "C3/C4 являются уровнями предварительного скрининга до экспертной проверки и не означают подтверждённого нарушения.", bold=True)

    # V — factual tables.
    new_section(doc)
    appendix_title(doc, "В", "Фактические таблицы анализа T3")
    titles = {
        "01_source_counts.csv": "В.1 — Распределение по источникам",
        "02_rating_distribution.csv": "В.2 — Распределение по рейтингу",
        "03_sentiment_proxy_distribution.csv": "В.3 — Распределение прокси-показателя тональности",
        "04_aspect_frequency.csv": "В.4 — Частота аспектов",
        "05_criticality_distribution.csv": "В.5 — Распределение критичности",
        "06_aspect_x_sentiment.csv": "В.6 — Аспект × прокси-показатель тональности",
        "07_aspect_x_criticality.csv": "В.7 — Аспект × критичность",
        "08_monthly_counts.csv": "В.8 — Распределение по месяцам",
    }
    for fname, title in titles.items():
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        rr = p.add_run(f"Таблица {title}"); rr.bold = True; rr.font.name = "Times New Roman"; rr.font.size = Pt(11)
        add_table(doc, csv_rows(TABLE_DIR / fname), font_size=9)
    add_body(doc, "Все значения перенесены непосредственно из закрытого T3 output без ручной корректировки.", italic=True)

    # G — charts.
    new_section(doc)
    appendix_title(doc, "Г", "Фактические диаграммы T3")
    chart_titles = [
        ("01_sources.png", "Рисунок Г.1 — Распределение корпуса по источникам"),
        ("02_ratings.png", "Рисунок Г.2 — Распределение рейтингов"),
        ("03_sentiment_proxy.png", "Рисунок Г.3 — Прокси-показатель тональности"),
        ("04_aspects.png", "Рисунок Г.4 — Частота аспектов"),
        ("05_criticality.png", "Рисунок Г.5 — Распределение критичности"),
        ("06_monthly.png", "Рисунок Г.6 — Динамика количества отзывов по месяцам"),
    ]
    for idx, (fname, caption) in enumerate(chart_titles):
        add_image(doc, CHART_DIR / fname, 14.5); add_caption(doc, caption)
        if idx in {1, 3}:
            doc.add_page_break()

    # D — P1-P8 matrix.
    new_section(doc, landscape=True)
    appendix_title(doc, "Д", "Итоговая матрица проблем и ограничений P1–P8")
    add_body(doc, "Матрица разграничивает подтверждённые публичной эмпирикой наблюдения и параметры, которые невозможно оценить по открытым данным.")
    p_rows = csv_rows(P1P8)
    add_table(doc, p_rows, font_size=7, widths_cm=[1.0, 2.6, 4.0, 3.5, 3.8, 5.8, 5.8])
    add_body(doc, "P4 и P5 сохраняют статус NOT ASSESSABLE INTERNALLY. Они используются как проектные требования и не являются доказанными недостатками внутренней системы «Перекрёстка».", bold=True)

    # E — TO-BE.
    new_section(doc)
    appendix_title(doc, "Е", "Целевая процессная модель TO-BE «Методики ЕГ»")
    tobe = dot_png("appendix_e_tobe", '''rankdir=TB;
A[label="Получение"] -> B[label="Регистрация + происхождение данных"] -> C[label="Валидация / дубли / ПДн"] -> D[label="Классификация"] -> E[label="Проверка C3/C4"] -> F[label="Маршрутизация"] -> G[label="Решение"];
G -> H[label="ответ"]; G -> I[label="мероприятие"]; H -> J[label="Контроль"]; I -> J; J -> K[label="Подтверждённый результат"]; K -> L[label="Закрытие / знания / аналитика"];
''')
    add_image(doc, tobe, 15.5); add_caption(doc, "Рисунок Е.1 — Целевой жизненный цикл работы с отзывом")
    add_bullets(doc, [
        "исходный канал и исходный идентификатор сохраняются при нормализации;",
        "тональность, аспекты и критичность рассматриваются независимо;",
        "C3/C4 требуют экспертной проверки;",
        "публичный ответ и внутреннее мероприятие являются разными сущностями;",
        "закрытие требует зафиксированного результата или обоснованного отсутствия дальнейшего действия;",
        "KPI рассчитываются из событий процесса, а не вводятся как заранее заявленный эффект.",
    ])

    # Zh — process RACI.
    new_section(doc, landscape=True)
    appendix_title(doc, "Ж", "Матрица ответственности RACI целевого процесса")
    add_body(doc, "Роли являются функциональными и не утверждают фактическую организационную структуру сети «Перекрёсток».")
    raci = [
        ["Этап", "PO", "FC", "AN", "SP", "CS", "LC", "RM", "QC"],
        ["Правила процесса", "A", "R", "R", "C", "C", "C", "I", "C"],
        ["Получение/регистрация", "I", "A/R", "C", "I", "R", "I", "I", "I"],
        ["Очистка/обезличивание", "I", "A", "R", "I", "C", "C", "I", "I"],
        ["Первичная классификация", "I", "A", "R", "C", "C", "I", "I", "I"],
        ["Проверка C3/C4", "I", "R", "R", "C", "C", "A/C", "A/C", "I"],
        ["Маршрутизация", "A", "R", "C", "C", "C", "C", "C", "I"],
        ["Решение C1/C2", "I", "A", "C", "R", "R", "C", "I", "I"],
        ["Решение C3", "I", "R", "C", "R", "C", "C", "A", "I"],
        ["Решение C4", "I", "R", "C", "C", "C", "C", "A", "I"],
        ["Публичный ответ", "I", "A", "C", "C", "R", "C*", "I", "I"],
        ["Внутреннее мероприятие", "I", "C", "I", "A/R", "C", "C", "C", "I"],
        ["Контроль исполнения", "A", "R", "C", "C", "C", "I", "I", "R"],
        ["Верификация результата", "A", "C", "C", "R", "C", "C", "I", "R"],
        ["Закрытие/повторное открытие", "A", "R", "I", "C", "C", "I", "I", "R"],
        ["Аналитика/KPI", "A", "C", "R", "I", "I", "I", "I", "C"],
    ]
    add_table(doc, raci, font_size=8, widths_cm=[6.0,2,2,2,2,2,2,2,2])
    add_body(doc, "R — выполняет; A — несёт итоговую ответственность; C — консультирует; I — информируется. C* — юридическая проверка для чувствительных случаев.")

    # I — classifier and escalation.
    new_section(doc)
    appendix_title(doc, "И", "Классификатор и дерево эскалации")
    add_table(doc, [["Измерение", "Содержание"],
                    ["Источник", "канал/площадка и реквизиты происхождения"],
                    ["Тональность", "POS / NEU / NEG / UNKNOWN"],
                    ["Аспекты", "многометочная классификация по базовому справочнику"],
                    ["Критичность", "C1–C4 независимо от тональности"],
                    ["Стадия", "статус целевого процесса"]], font_size=10, widths_cm=[4.5,11])
    add_bullets(doc, ["безопасность/здоровье", "персональные данные/конфиденциальность", "существенный финансовый ущерб", "права потребителей/претензия", "массовая повторяемость", "высокая публичность/медиариск", "повторный нерешённый случай", "отсутствие полномочий на стандартном уровне"])
    decision = dot_png("appendix_i_decision", '''rankdir=TB;
A[label="Новый отзыв"] -> B[label="Классификация"] -> C[label="Критерий критичности?",shape=diamond]; C -> D[label="Предварительный C3/C4", labeldistance=1]; D -> E[label="Экспертная проверка"]; E -> F[label="Высокий риск подтверждён?",shape=diamond]; F -> G[label="Эскалация", label="да"]; F -> H[label="Изменить уровень с основанием", label="нет"]; C -> I[label="Стандартная маршрутизация", label="нет"]; G -> J[label="Решение"]; H -> J; I -> J; J -> K[label="Ответ и/или мероприятие"] -> L[label="Контроль"] -> M[label="Результат подтверждён?",shape=diamond]; M -> J[label="нет"]; M -> N[label="Закрытие",label="да"];
''')
    add_image(doc, decision, 15.0); add_caption(doc, "Рисунок И.1 — Логика проверки и эскалации")

    # K — architecture.
    new_section(doc)
    appendix_title(doc, "К", "Функциональная и информационная архитектура")
    arch = dot_png("appendix_k_arch", '''rankdir=LR;
A[label="Каналы"] -> B[label="Регистрация"] -> C[label="Качество данных"] -> D[label="Классификация"] -> E[label="Решение/маршрутизация"]; E -> F[label="Коммуникация"]; E -> G[label="Мероприятия"]; F -> H[label="Контроль результата"]; G -> H; H -> I[label="Аналитика/KPI"] -> J[label="База знаний"] -> D;
''')
    add_image(doc, arch, 15.5); add_caption(doc, "Рисунок К.1 — Функциональная архитектура")
    entities = ["SOURCE", "REVIEW_CASE", "CLASSIFICATION", "ASPECT", "DECISION", "RESPONSE", "ACTION", "CONTROL_EVENT", "OUTCOME", "ROLE", "AUDIT_EVENT"]
    add_table(doc, [["Сущность", "Назначение"]] + [[e, {
        "SOURCE":"источник и реквизиты происхождения", "REVIEW_CASE":"исходная карточка отзыва", "CLASSIFICATION":"тональность, аспекты, C-level и метод", "ASPECT":"справочник аспектов", "DECISION":"решение и основание", "RESPONSE":"коммуникационный ответ", "ACTION":"внутреннее мероприятие", "CONTROL_EVENT":"событие контроля", "OUTCOME":"подтверждённый результат", "ROLE":"функциональная роль", "AUDIT_EVENT":"история существенных изменений"}[e]] for e in entities], font_size=9, widths_cm=[5,10.5])
    add_body(doc, "Архитектура платформенно-независима и не является доказательством реализованной программной системы.", bold=True)

    # L — S01-S10 conceptual screens.
    new_section(doc)
    appendix_title(doc, "Л", "Концептуальные проектные представления S01–S10")
    add_body(doc, "Представления S01–S10 визуализируют информационную архитектуру будущего решения. Они не являются скриншотами готового программного обеспечения.", bold=True)
    screens = [
        ("S01", "Дашборд", "новые отзывы; открытые C3/C4; SLA; динамика; аспекты; просроченные мероприятия; доля публичных ответов отдельно от подтверждённого решения"),
        ("S02", "Список отзывов", "фильтры по периоду, источнику, аспекту, C-level, статусу, сроку; основные поля кейса"),
        ("S03", "Карточка отзыва", "исходный текст; происхождение; аналитическая интерпретация; C1–C4; история; решения/ответы/мероприятия/результат"),
        ("S04", "Импорт/получение", "источник; сопоставление полей; принятые/исключённые/дубли; качество данных; обезличивание"),
        ("S05", "Решения", "реестр решений, тип, основание, ответственная роль, критичность, срок, статус"),
        ("S06", "Карточка решения", "основание; ответ; мероприятие; владелец; срок; эскалация; подтверждённый результат; история"),
        ("S07", "Контроль", "просрочки; C3/C4 без проверки; ответ без завершённого мероприятия; результат без проверки; повторно открытые случаи"),
        ("S08", "Аналитика", "динамика; каналы; аспекты; критичность; ответ vs подтверждённое решение; повторяемость; переход к исходному кейсу"),
        ("S09", "Справочники", "аспекты; критерии критичности; C1–C4; роли; типы решений/мероприятий/результатов; проектные SLA"),
        ("S10", "Настройки/методика", "версия кодировочной схемы; обезличивание; сопоставление каналов; RACI; KPI; правила доказательности и аудита"),
    ]
    for sid, name, blocks in screens:
        add_table(doc, [[f"{sid} — {name}", "Концептуальное содержание"], ["Назначение", blocks], ["Статус", "проектное представление; не evidence программной реализации"]], font_size=9, widths_cm=[4.5,11])

    # M — KPI.
    new_section(doc, landscape=True)
    appendix_title(doc, "М", "Система KPI и формулы")
    kpis = [
        ("KPI-01", "Количество зарегистрированных отзывов", "count(unique review cases)"),
        ("KPI-02", "Полнота происхождения данных", "records_with_required_source_fields / all_records × 100%"),
        ("KPI-03", "Доля дублей", "duplicates / raw_records × 100%"),
        ("KPI-04", "Медианное время до первичной обработки", "median(triage_at − registered_at)"),
        ("KPI-05", "Медианное время до маршрутизации", "median(routed_at − registered_at)"),
        ("KPI-06", "Соблюдение SLA", "within_SLA / applicable_SLA_cases × 100%"),
        ("KPI-07", "Доля экспертно проверенных C3/C4", "verified_C3_C4 / screened_C3_C4 × 100%"),
        ("KPI-08", "Доля реклассификации C3/C4", "changed_after_review / reviewed_C3_C4 × 100%"),
        ("KPI-09", "Доля отзывов с 2+ аспектами", "reviews_2plus_aspects / classified_reviews × 100%"),
        ("KPI-10", "Доля публичных ответов", "public_response / response_applicable × 100%"),
        ("KPI-11", "Доля случаев, требующих мероприятия", "internal_action_required / processed × 100%"),
        ("KPI-12", "Доля подтверждённо решённых случаев", "verified_resolution / resolution_required × 100%"),
        ("KPI-13", "Доля повторного открытия", "reopened / closed × 100%"),
        ("KPI-14", "Доля повторяющихся проблем", "repeat_problem_cases / analyzed_cases × 100%"),
        ("KPI-15", "Завершение корректирующих мероприятий", "completed_verified_actions / due_actions × 100%"),
        ("KPI-16", "Динамика аспекта", "количество/доля аспекта по сопоставимым периодам"),
    ]
    add_table(doc, [["ID", "Показатель", "Формула/правило"]] + [list(x) for x in kpis], font_size=8, widths_cm=[2.0,9.0,15.0])
    add_body(doc, "KPI являются проектной системой будущего измерения и не представляют фактическую эффективность сети «Перекрёсток».", bold=True)

    # N — implementation E0-E8.
    new_section(doc)
    appendix_title(doc, "Н", "План внедрения E0–E8")
    impl = dot_png("appendix_n_impl", '''rankdir=LR;
E0[label="E0\nРешение"] -> E1[label="E1\nПодготовка"] -> E2[label="E2\nАдаптация"] -> E3[label="E3\nОбучение"] -> E4[label="E4\nПериод «до»"] -> E5[label="E5\nПилот"] -> E6[label="E6\nИзмерение «после»"] -> E7[label="E7\nОценка"] -> D[label="Решение",shape=diamond]; D -> E8[label="E8\nМасштабирование"]; D -> C[label="Корректировать"]; D -> X[label="Остановить"];
''')
    add_image(doc, impl, 15.5); add_caption(doc, "Рисунок Н.1 — Последовательность будущего внедрения")
    stages = [
        ("E0", "решение о пилоте", "владелец, область, доступные каналы, ограничения"),
        ("E1", "подготовка", "инвентаризация каналов, базовый период, источники данных"),
        ("E2", "адаптация", "аспекты, критерии критичности, SLA, роли, KPI"),
        ("E3", "обучение и пробный прогон", "разбор спорных случаев и маршрутизации"),
        ("E4", "период «до»", "фиксация фактической исходной линии"),
        ("E5", "пилот", "работа по TO-BE на заранее выбранных границах"),
        ("E6", "измерение «после»", "расчёт по заранее утверждённым формулам"),
        ("E7", "оценка", "сопоставимость, риски, KPI, guardrails"),
        ("E8", "масштабирование", "только после решения «масштабировать»"),
    ]
    add_table(doc, [["Этап", "Содержание", "Ключевой результат"]] + [list(x) for x in stages], font_size=9, widths_cm=[2.2,5.0,8.3])
    add_body(doc, "ВКР содержит план будущего внедрения; фактический пилот не проводился.", bold=True)

    # P — implementation RACI & training.
    new_section(doc, landscape=True)
    appendix_title(doc, "П", "RACI внедрения и программа обучения")
    iraci = [
        ["Работа", "SP", "PO", "PC", "PR", "CS", "BO", "AN", "LC", "SA"],
        ["Утвердить границы пилота", "A", "R", "C", "C", "C", "C", "C", "C", "I"],
        ["Утвердить методику/KPI", "I", "A", "R", "C", "C", "C", "R", "C", "I"],
        ["Утвердить ПДн/доступ", "I", "C", "R", "I", "I", "I", "I", "A/R", "C"],
        ["Настроить реестр/систему", "I", "C", "A", "I", "C", "C", "C", "C", "R"],
        ["Провести обучение", "I", "A", "R", "R", "R", "R", "C", "C", "C"],
        ["Сформировать период «до»", "I", "A", "R", "C", "C", "C", "R", "C", "C"],
        ["Вести пилот", "I", "A", "R", "R", "R", "R", "C", "C", "C"],
        ["Проверка C3/C4", "I", "A", "R", "C", "C", "C", "C", "R/C", "I"],
        ["Сформировать период «после»", "I", "A", "R", "C", "C", "C", "R", "C", "C"],
        ["Оценить результат", "C", "A", "R", "C", "C", "C", "R", "C", "I"],
        ["Решение о масштабировании", "A", "R", "C", "C", "C", "C", "C", "C", "I"],
    ]
    add_table(doc, iraci, font_size=7, widths_cm=[7.0,1.7,1.7,1.7,1.7,1.7,1.7,1.7,1.7,1.7])
    new_section(doc)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rr = p.add_run("Программа обучения"); rr.bold=True; rr.font.name="Times New Roman"; rr.font.size=Pt(14)
    add_table(doc, [
        ["Модуль", "Содержание"],
        ["1. Логика методики", "отзыв, решение, ответ, мероприятие, результат; происхождение данных; жизненный цикл"],
        ["2. Классификация", "тональность; многометочные аспекты; C1–C4; спорные случаи"],
        ["3. Маршрутизация", "владелец; SLA; проверка C3/C4; юридические/ПДн критерии"],
        ["4. Коммуникация", "публичный ответ и его связь с внутренним мероприятием"],
        ["5. Контроль результата", "контрольное событие; подтверждённый результат; повторное открытие"],
        ["6. KPI и доказательность", "формулы, воспроизводимость, запрет менять правила после просмотра результата"],
    ], font_size=10, widths_cm=[5.0,10.5])

    # R — risks.
    new_section(doc, landscape=True)
    appendix_title(doc, "Р", "Карта рисков R1–R20 и stop-risks")
    risks = [
        ("R1","Неполный охват каналов","средняя","высокое","реестр каналов, KPI полноты происхождения, поэтапное подключение"),
        ("R2","Дубли между источниками","средняя","среднее","правила дедупликации, source ID, отпечаток текст/дата"),
        ("R3","Несогласованная классификация","высокая на старте","высокое","кодировочная схема, обучение, аудит реклассификации"),
        ("R4","Ложная C3/C4 эскалация","средняя","среднее","предварительный скрининг + экспертная проверка"),
        ("R5","Пропуск критичного случая","низкая/средняя","критическое","независимые критерии, ручная проверка, контроль эскалации"),
        ("R6","Подмена результата публичным ответом","средняя","высокое","раздельные RESPONSE/ACTION/OUTCOME"),
        ("R7","Нарушение ПДн","низкая при контроле","критическое","минимизация, разграничение доступа, юридическая проверка"),
        ("R8","Обход регламента участниками","средняя","высокое","обучение, владелец, журнал изменений, исключения"),
        ("R9","Формальный SLA без контроля","средняя","высокое","события времени, просрочки, эскалация владельца"),
        ("R10","Перегрузка первой линии","средняя","высокое","контроль нагрузки, C1–C4, ограничение пилота"),
        ("R11","Слишком сложный классификатор","средняя","среднее","минимальный набор полей, анализ OTHER, упрощение"),
        ("R12","Изменение кодировочной схемы в пилоте","средняя","высокое","заморозка версии, change log"),
        ("R13","Несопоставимость «до/после»","средняя","критическое","предварительная фиксация, одинаковые границы и формулы"),
        ("R14","Сезонность/внешние события","средняя","высокое","сопоставимые окна, стратификация, журнал контекста"),
        ("R15","Недостаточный объём «после»","средняя","высокое","правило минимальной выборки до старта"),
        ("R16","Изменение канала/API","средняя","среднее","абстракция источника, ручной резервный путь"),
        ("R17","Нет владельца мероприятия","средняя","высокое","RACI, обязательный владелец, эскалация просрочки"),
        ("R18","KPI становятся самоцелью","средняя","высокое","сбалансированный набор и качественные проверки"),
        ("R19","Порог успеха меняется после результатов","низкая при управлении","критическое","предварительно зафиксированные пороги и версия"),
        ("R20","Преждевременное масштабирование","средняя","высокое","формальная контрольная точка и backlog корректирующих мер"),
    ]
    add_table(doc, [["ID","Риск","Вероятность","Влияние","Меры снижения"]] + [list(x) for x in risks], font_size=7, widths_cm=[1.5,7.0,3.2,3.0,12.0])
    add_body(doc, "Stop-risks: существенный инцидент ПДн/безопасности; систематическая потеря происхождения данных; невоспроизводимость KPI; разные формулы «до/после»; массовое использование неутверждённой версии классификатора; отсутствие владельца C4.", bold=True)

    # S — BEFORE/AFTER protocol.
    new_section(doc, landscape=True)
    appendix_title(doc, "С", "Протокол будущего сравнения «до/после»")
    add_body(doc, "Методика фиксирует дизайн будущего измерения. Фактические значения AFTER в ВКР отсутствуют.")
    ba = [
        ("K1","Полнота происхождения","cases_with_required_provenance / all_cases × 100%"),
        ("K2","Полнота классификации","cases_with_required_classification / all_cases × 100%"),
        ("K3","Медианное время до первичной обработки","median(triage_timestamp − registered_timestamp)"),
        ("K4","Медианное время до маршрутизации","median(routed_timestamp − registered_timestamp)"),
        ("K5","Соблюдение SLA","within_target / applicable_SLA × 100%"),
        ("K6","Проверка C3/C4","verified_C3_C4 / screened_C3_C4 × 100%"),
        ("K7","Реклассификация","reclassified / classified × 100%"),
        ("K8","Публичный ответ","public_response / response_required × 100%"),
        ("K9","Охват мероприятиями","action_created / internal_action_required × 100%"),
        ("K10","Подтверждённое решение","verified_positive_outcome / resolution_required × 100%"),
        ("K11","Повторное открытие","reopened / closed × 100%"),
        ("K12","Повторяющиеся проблемы","repeat_problem / problem_cases × 100%"),
        ("K13","Завершение корректирующих мер","closed_with_evidence / due_actions × 100%"),
        ("K14","Полнота прослеживаемости","review_decision_links / applicable_cases × 100%"),
        ("K15","Исключения качества данных","data_quality_exception / all_cases × 100%"),
    ]
    add_table(doc, [["ID","Метрика","Формула"]] + [list(x) for x in ba], font_size=8, widths_cm=[2.0,9.0,15.0])
    add_body(doc, "Для долей: Δpp = AFTER% − BEFORE%. Для временных метрик: Δtime = AFTERmedian − BEFOREmedian. BEFORE и AFTER должны иметь сопоставимые границы, каналы, правила включения/исключения, формулы, сезонность и правила обработки дублей.")
    add_body(doc, "Запрещено удалять показатель из итогового отчёта только потому, что он не улучшился (anti-cherry-picking).", bold=True)

    # T — decision gate.
    new_section(doc)
    appendix_title(doc, "Т", "Контрольная точка принятия решения после будущего пилота")
    add_body(doc, "Решение после пилота принимается только на воспроизводимой базе и не заменяется искусственным интегральным баллом.")
    add_table(doc, [
        ["Блок", "Условие"],
        ["Integrity gates", "воспроизводимость; сопоставимость формул; закрытые существенные ПДн/безопасностные инциденты; владелец C4; наличие evidence для ключевых итогов"],
        ["Guardrails", "privacy/security; отсутствие пропуска C3/C4; подтверждённое решение; качество данных; приемлемая нагрузка; контролируемая реклассификация/повторное открытие"],
        ["МАСШТАБИРОВАТЬ", "обязательные gates и guardrails PASS; заранее утверждённые критерии выполнены; нет открытого high-risk блокера"],
        ["КОРРЕКТИРОВАТЬ", "данные воспроизводимы, stop-risk отсутствует, но есть исправимые проблемы классификатора, RACI, SLA, обучения или процесса"],
        ["ОСТАНОВИТЬ", "целостность исследования невосстановима, неприемлемый privacy/safety риск, критические guardrails ухудшаются или пилот не позволяет надёжный вывод"],
    ], font_size=10, widths_cm=[4.2,11.3])
    add_table(doc, [
        ["Поле предварительной регистрации", "Значение до старта пилота"],
        ["Pilot ID", "задаётся"], ["Версия методики", "задаётся"], ["Границы", "задаются"], ["Период «до»", "задаётся"], ["Период пилота/«после»", "задаётся"], ["Каналы", "задаются"], ["Правила включения/исключения", "задаются"], ["Список KPI и формулы", "фиксируются"], ["Пороги успеха", "утверждаются до результатов"], ["Guardrails", "утверждаются"], ["Stop-risks", "утверждаются"], ["Лицо, принимающее решение", "назначается"],
    ], font_size=9, widths_cm=[7.0,8.5])
    add_body(doc, "Пункты результатов пилота и AFTER не заполняются до фактического проведения пилота.", bold=True)

    # Metadata and save.
    doc.core_properties.title = "Приложения А–Т к ВКР — Методика ЕГ"
    doc.core_properties.subject = "T7.1 pre-defense appendix package"
    doc.save(DOCX)
    print(DOCX)


if __name__ == "__main__":
    build()
