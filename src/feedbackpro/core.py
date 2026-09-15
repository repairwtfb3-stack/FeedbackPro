from __future__ import annotations

import csv
import hashlib
import json
import re
import sqlite3
from contextlib import contextmanager
from dataclasses import dataclass, field, replace
from datetime import date, datetime
from enum import StrEnum
from pathlib import Path
from typing import Any, Iterable, Iterator, Mapping

ANALYSIS_VERSION = "a2-rules-1.0"

class Sentiment(StrEnum):
    POSITIVE = "positive"
    NEUTRAL = "neutral"
    NEGATIVE = "negative"

class Criticality(StrEnum):
    LOW = "low"
    MEDIUM = "medium"
    HIGH = "high"
    CRITICAL = "critical"

class DecisionStatus(StrEnum):
    CREATED = "created"
    PLANNED = "planned"
    IN_PROGRESS = "in_progress"
    COMPLETED = "completed"
    VERIFIED = "verified"

class ControlStatus(StrEnum):
    PLANNED = "planned"
    IN_PROGRESS = "in_progress"
    COMPLETED = "completed"
    VERIFIED = "verified"

@dataclass(frozen=True, slots=True)
class Review:
    id: int | None
    source: str
    channel: str
    review_date: date | None
    author: str | None
    rating: float | None
    text: str
    normalized_text: str
    fingerprint: str
    external_id: str | None = None
    source_file: str | None = None
    imported_at: datetime = field(default_factory=datetime.utcnow)

@dataclass(frozen=True, slots=True)
class AnalysisResult:
    review_id: int | None
    sentiment: Sentiment
    sentiment_score: float
    confidence: float
    aspects: tuple[str, ...]
    criticality: Criticality
    explanation: str
    analysis_version: str
    analyzed_at: datetime = field(default_factory=datetime.utcnow)

_SPACE_RE = re.compile(r"\s+")
_TOKEN_RE = re.compile(r"[A-Za-zА-Яа-яЁё0-9-]+", re.UNICODE)
POSITIVE = {"отлично", "отличный", "хорошо", "хороший", "быстро", "удобно", "доволен", "довольна", "рекомендую", "спасибо", "качественный", "вежливый", "супер", "прекрасно", "понравилось", "идеально"}
NEGATIVE = {"плохо", "плохой", "ужасно", "ужасный", "медленно", "грубо", "сломался", "сломано", "брак", "просрочка", "задержка", "обман", "ошибка", "верните", "жалоба", "недоволен", "недовольна", "некачественный", "хамство"}
INTENSIFIERS = {"очень", "крайне", "совсем", "абсолютно", "дважды", "повторно"}
NEGATIONS = {"не", "нет", "никогда", "ни"}
ASPECT_RULES = {
    "quality": {"качество", "брак", "сломался", "сломано", "дефект", "товар"},
    "price": {"цена", "дорого", "дешево", "стоимость", "скидка"},
    "delivery": {"доставка", "курьер", "привезли", "задержка", "срок"},
    "service": {"сервис", "обслуживание", "менеджер", "консультант"},
    "staff": {"персонал", "сотрудник", "кассир", "хамство", "грубо"},
    "support": {"поддержка", "оператор", "чат", "звонок", "ответ"},
    "return": {"возврат", "верните", "гарантия", "обмен"},
    "digital": {"сайт", "приложение", "личный", "кабинет", "оплата"},
    "availability": {"наличие", "нет", "закончился", "склад"},
}
CRITICAL_MARKERS = {"деньги", "списали", "мошенничество", "опасно", "травма", "суд", "прокуратура", "роспотребнадзор", "утечка", "персональные", "данные"}
HIGH_MARKERS = {"обман", "повторно", "дважды", "жалоба", "брак", "верните"}

def normalize_text(value: str) -> str:
    return _SPACE_RE.sub(" ", (value or "").strip())

def _parse_date(value: Any) -> date | None:
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    raw = str(value).strip()
    for fmt in ("%Y-%m-%d", "%d.%m.%Y", "%d/%m/%Y", "%Y/%m/%d"):
        try:
            return datetime.strptime(raw, fmt).date()
        except ValueError:
            pass
    raise ValueError(f"Unsupported review date: {raw}")

def _parse_rating(value: Any) -> float | None:
    if value in (None, ""):
        return None
    rating = float(str(value).replace(",", "."))
    if not 0 <= rating <= 5:
        raise ValueError("rating must be between 0 and 5")
    return rating

def make_fingerprint(source: str, external_id: str | None, review_date: date | None, text: str) -> str:
    raw = f"id|{source.casefold()}|{external_id.strip().casefold()}" if external_id else f"text|{source.casefold()}|{review_date or ''}|{text.casefold()}"
    return hashlib.sha256(raw.encode("utf-8")).hexdigest()

def normalize_record(record: Mapping[str, Any], *, source_file: str | None = None) -> Review:
    source = normalize_text(str(record.get("source") or "unknown"))
    channel = normalize_text(str(record.get("channel") or source))
    text = normalize_text(str(record.get("text") or ""))
    if not text:
        raise ValueError("review text is required")
    review_date = _parse_date(record.get("date") or record.get("review_date"))
    ext = record.get("external_id") or record.get("id")
    external_id = normalize_text(str(ext)) if ext not in (None, "") else None
    author_raw = record.get("author")
    author = normalize_text(str(author_raw)) if author_raw not in (None, "") else None
    rating = _parse_rating(record.get("rating"))
    return Review(id=None, source=source, channel=channel, review_date=review_date, author=author, rating=rating, text=text, normalized_text=text, fingerprint=make_fingerprint(source, external_id, review_date, text), external_id=external_id, source_file=source_file)

def _tokens(text: str) -> list[str]:
    return [x.casefold() for x in _TOKEN_RE.findall(text)]

class AnalysisPipeline:
    def analyze(self, review: Review) -> AnalysisResult:
        tokens = _tokens(review.normalized_text)
        score = 0.0
        hits: list[str] = []
        for i, token in enumerate(tokens):
            weight = 1.5 if i and tokens[i - 1] in INTENSIFIERS else 1.0
            negated = bool(i and tokens[i - 1] in NEGATIONS)
            if token in POSITIVE:
                score += -weight if negated else weight
                hits.append(("not " if negated else "") + token)
            elif token in NEGATIVE:
                score += weight if negated else -weight
                hits.append(("not " if negated else "") + token)
        if review.rating is not None:
            if review.rating <= 2:
                score -= 0.75
            elif review.rating >= 4:
                score += 0.75
        sentiment = Sentiment.POSITIVE if score > 0.35 else Sentiment.NEGATIVE if score < -0.35 else Sentiment.NEUTRAL
        confidence = min(0.98, 0.50 + 0.10 * abs(score) + 0.03 * len(hits))
        token_set = set(tokens)
        aspects = tuple([k for k, words in ASPECT_RULES.items() if token_set & words] or ["other"])
        risk = 2 if sentiment is Sentiment.NEGATIVE else 0
        critical_hits = len(token_set & CRITICAL_MARKERS)
        if critical_hits:
            risk += 4 + critical_hits
        if token_set & HIGH_MARKERS:
            risk += 2
        if review.rating is not None and review.rating <= 1:
            risk += 1
        if {"return", "support"} & set(aspects):
            risk += 1
        criticality = Criticality.CRITICAL if risk >= 7 else Criticality.HIGH if risk >= 4 else Criticality.MEDIUM if risk >= 2 else Criticality.LOW
        explanation = f"rule_score={score:.2f}; markers={','.join(hits[:8]) or 'none'}; aspects={','.join(aspects)}; criticality={criticality.value}; version={ANALYSIS_VERSION}"
        return AnalysisResult(review_id=review.id, sentiment=sentiment, sentiment_score=round(score, 3), confidence=round(confidence, 3), aspects=aspects, criticality=criticality, explanation=explanation, analysis_version=ANALYSIS_VERSION)

MIGRATION_V1 = """
CREATE TABLE IF NOT EXISTS schema_migrations(version INTEGER PRIMARY KEY, applied_at TEXT NOT NULL);
CREATE TABLE IF NOT EXISTS reviews(
 id INTEGER PRIMARY KEY AUTOINCREMENT,
 source TEXT NOT NULL, channel TEXT NOT NULL, review_date TEXT NULL, author TEXT NULL,
 rating REAL NULL, text TEXT NOT NULL, normalized_text TEXT NOT NULL,
 fingerprint TEXT NOT NULL UNIQUE, external_id TEXT NULL, source_file TEXT NULL,
 imported_at TEXT NOT NULL
);
CREATE TABLE IF NOT EXISTS settings(key TEXT PRIMARY KEY, value TEXT NOT NULL);
"""
MIGRATION_V2 = """
CREATE TABLE IF NOT EXISTS analysis_results(
 review_id INTEGER PRIMARY KEY, sentiment TEXT NOT NULL, sentiment_score REAL NOT NULL,
 confidence REAL NOT NULL, aspects_json TEXT NOT NULL, criticality TEXT NOT NULL,
 explanation TEXT NOT NULL, analysis_version TEXT NOT NULL, analyzed_at TEXT NOT NULL,
 FOREIGN KEY(review_id) REFERENCES reviews(id) ON DELETE CASCADE
);
CREATE TABLE IF NOT EXISTS topics(code TEXT PRIMARY KEY, title TEXT NOT NULL, active INTEGER NOT NULL DEFAULT 1);
CREATE TABLE IF NOT EXISTS decisions(
 id INTEGER PRIMARY KEY AUTOINCREMENT, title TEXT NOT NULL, description TEXT NOT NULL,
 priority TEXT NOT NULL, status TEXT NOT NULL, due_date TEXT NULL, created_at TEXT NOT NULL
);
CREATE TABLE IF NOT EXISTS decision_reviews(
 decision_id INTEGER NOT NULL, review_id INTEGER NOT NULL,
 PRIMARY KEY(decision_id,review_id),
 FOREIGN KEY(decision_id) REFERENCES decisions(id) ON DELETE CASCADE,
 FOREIGN KEY(review_id) REFERENCES reviews(id) ON DELETE CASCADE
);
CREATE TABLE IF NOT EXISTS control_items(
 id INTEGER PRIMARY KEY AUTOINCREMENT, decision_id INTEGER NOT NULL UNIQUE,
 status TEXT NOT NULL, due_date TEXT NULL, outcome TEXT NULL, verified_at TEXT NULL,
 FOREIGN KEY(decision_id) REFERENCES decisions(id) ON DELETE CASCADE
);
"""
DEFAULT_TOPICS = {
    "quality": "Качество товара", "price": "Цена", "delivery": "Доставка", "service": "Сервис", "staff": "Персонал", "support": "Поддержка", "return": "Возврат и обмен", "digital": "Сайт/приложение", "availability": "Наличие", "other": "Прочее",
}

class Database:
    def __init__(self, path: str | Path = "feedbackpro.db") -> None:
        self.path = str(path)
        self.conn = sqlite3.connect(self.path)
        self.conn.row_factory = sqlite3.Row
        self.conn.execute("PRAGMA foreign_keys=ON")
    def close(self) -> None:
        self.conn.close()
    def __enter__(self) -> "Database":
        return self
    def __exit__(self, *exc: object) -> None:
        self.close()
    @contextmanager
    def transaction(self) -> Iterator[sqlite3.Connection]:
        try:
            self.conn.execute("BEGIN")
            yield self.conn
            self.conn.commit()
        except Exception:
            self.conn.rollback()
            raise
    def current_version(self) -> int:
        if self.conn.execute("SELECT 1 FROM sqlite_master WHERE type='table' AND name='schema_migrations'").fetchone() is None:
            return 0
        row = self.conn.execute("SELECT MAX(version) AS v FROM schema_migrations").fetchone()
        return int(row["v"] or 0)
    def migrate(self, target: int = 2) -> int:
        if target not in (1, 2):
            raise ValueError("supported schema targets are 1 and 2")
        current = self.current_version()
        if current < 1 and target >= 1:
            self.conn.executescript(MIGRATION_V1)
            self.conn.execute("INSERT OR IGNORE INTO schema_migrations VALUES(1,?)", (datetime.utcnow().isoformat(),))
            self.conn.commit(); current = 1
        if current < 2 and target >= 2:
            self.conn.executescript(MIGRATION_V2)
            self.conn.execute("INSERT OR IGNORE INTO schema_migrations VALUES(2,?)", (datetime.utcnow().isoformat(),))
            self.conn.executemany("INSERT OR IGNORE INTO topics(code,title,active) VALUES(?,?,1)", DEFAULT_TOPICS.items())
            self.conn.commit(); current = 2
        return current
    def add_review(self, review: Review) -> tuple[Review, bool]:
        try:
            cur = self.conn.execute("""INSERT INTO reviews(source,channel,review_date,author,rating,text,normalized_text,fingerprint,external_id,source_file,imported_at) VALUES(?,?,?,?,?,?,?,?,?,?,?)""", (review.source, review.channel, review.review_date.isoformat() if review.review_date else None, review.author, review.rating, review.text, review.normalized_text, review.fingerprint, review.external_id, review.source_file, review.imported_at.isoformat()))
            self.conn.commit(); return replace(review, id=int(cur.lastrowid)), True
        except sqlite3.IntegrityError:
            self.conn.rollback()
            row = self.conn.execute("SELECT * FROM reviews WHERE fingerprint=?", (review.fingerprint,)).fetchone()
            return self._review(row), False
    def _review(self, row: sqlite3.Row) -> Review:
        return Review(id=row["id"], source=row["source"], channel=row["channel"], review_date=date.fromisoformat(row["review_date"]) if row["review_date"] else None, author=row["author"], rating=row["rating"], text=row["text"], normalized_text=row["normalized_text"], fingerprint=row["fingerprint"], external_id=row["external_id"], source_file=row["source_file"], imported_at=datetime.fromisoformat(row["imported_at"]))
    def list_reviews(self) -> list[Review]:
        return [self._review(r) for r in self.conn.execute("SELECT * FROM reviews ORDER BY id").fetchall()]
    def get_review(self, review_id: int) -> Review:
        row = self.conn.execute("SELECT * FROM reviews WHERE id=?", (review_id,)).fetchone()
        if row is None: raise KeyError(review_id)
        return self._review(row)
    def save_analysis(self, result: AnalysisResult) -> None:
        if result.review_id is None: raise ValueError("persist review before analysis")
        self.conn.execute("""INSERT INTO analysis_results VALUES(?,?,?,?,?,?,?,?,?) ON CONFLICT(review_id) DO UPDATE SET sentiment=excluded.sentiment,sentiment_score=excluded.sentiment_score,confidence=excluded.confidence,aspects_json=excluded.aspects_json,criticality=excluded.criticality,explanation=excluded.explanation,analysis_version=excluded.analysis_version,analyzed_at=excluded.analyzed_at""", (result.review_id, result.sentiment.value, result.sentiment_score, result.confidence, json.dumps(result.aspects, ensure_ascii=False), result.criticality.value, result.explanation, result.analysis_version, result.analyzed_at.isoformat()))
        self.conn.commit()
    def get_analysis(self, review_id: int) -> AnalysisResult | None:
        r = self.conn.execute("SELECT * FROM analysis_results WHERE review_id=?", (review_id,)).fetchone()
        if r is None: return None
        return AnalysisResult(review_id=r["review_id"], sentiment=Sentiment(r["sentiment"]), sentiment_score=r["sentiment_score"], confidence=r["confidence"], aspects=tuple(json.loads(r["aspects_json"])), criticality=Criticality(r["criticality"]), explanation=r["explanation"], analysis_version=r["analysis_version"], analyzed_at=datetime.fromisoformat(r["analyzed_at"]))
    def review_records(self) -> list[dict[str, Any]]:
        rows = self.conn.execute("""SELECT r.*,a.sentiment,a.sentiment_score,a.confidence,a.aspects_json,a.criticality,a.analysis_version FROM reviews r LEFT JOIN analysis_results a ON a.review_id=r.id ORDER BY r.id DESC""").fetchall()
        return [{"id":r["id"],"source":r["source"],"channel":r["channel"],"review_date":r["review_date"],"author":r["author"],"rating":r["rating"],"text":r["text"],"sentiment":r["sentiment"],"sentiment_score":r["sentiment_score"],"confidence":r["confidence"],"aspects":json.loads(r["aspects_json"]) if r["aspects_json"] else [],"criticality":r["criticality"],"analysis_version":r["analysis_version"]} for r in rows]
    def create_decision(self, title: str, description: str, priority: str, review_ids: Iterable[int], due_date: date | None = None) -> int:
        ids = tuple(dict.fromkeys(int(x) for x in review_ids))
        if not ids: raise ValueError("decision must reference review")
        with self.transaction() as cx:
            for rid in ids:
                if cx.execute("SELECT 1 FROM reviews WHERE id=?", (rid,)).fetchone() is None: raise KeyError(rid)
            cur = cx.execute("INSERT INTO decisions(title,description,priority,status,due_date,created_at) VALUES(?,?,?,?,?,?)", (title.strip(), description.strip(), priority.strip() or "medium", DecisionStatus.CREATED.value, due_date.isoformat() if due_date else None, datetime.utcnow().isoformat()))
            did = int(cur.lastrowid)
            cx.executemany("INSERT INTO decision_reviews VALUES(?,?)", [(did, rid) for rid in ids])
        return did
    def decision_records(self) -> list[dict[str, Any]]:
        rows = self.conn.execute("""SELECT d.*,GROUP_CONCAT(dr.review_id) review_ids,c.status control_status,c.outcome control_outcome FROM decisions d LEFT JOIN decision_reviews dr ON dr.decision_id=d.id LEFT JOIN control_items c ON c.decision_id=d.id GROUP BY d.id ORDER BY d.id DESC""").fetchall()
        return [{"id":r["id"],"title":r["title"],"description":r["description"],"priority":r["priority"],"status":r["status"],"due_date":r["due_date"],"review_ids":[int(x) for x in (r["review_ids"] or "").split(",") if x],"control_status":r["control_status"],"control_outcome":r["control_outcome"]} for r in rows]
    def set_control(self, decision_id: int, status: ControlStatus, due_date: date | None = None, outcome: str | None = None) -> int:
        if self.conn.execute("SELECT 1 FROM decisions WHERE id=?", (decision_id,)).fetchone() is None: raise KeyError(decision_id)
        verified = datetime.utcnow().isoformat() if status is ControlStatus.VERIFIED else None
        self.conn.execute("""INSERT INTO control_items(decision_id,status,due_date,outcome,verified_at) VALUES(?,?,?,?,?) ON CONFLICT(decision_id) DO UPDATE SET status=excluded.status,due_date=excluded.due_date,outcome=excluded.outcome,verified_at=excluded.verified_at""", (decision_id, status.value, due_date.isoformat() if due_date else None, outcome, verified))
        self.conn.commit(); return int(self.conn.execute("SELECT id FROM control_items WHERE decision_id=?", (decision_id,)).fetchone()["id"])
    def control_records(self) -> list[dict[str, Any]]:
        return [dict(r) for r in self.conn.execute("""SELECT c.*,d.title decision_title FROM control_items c JOIN decisions d ON d.id=c.decision_id ORDER BY c.id DESC""").fetchall()]
    def topics(self) -> list[dict[str, Any]]:
        return [dict(r) for r in self.conn.execute("SELECT * FROM topics ORDER BY title").fetchall()]
    def set_topic(self, code: str, title: str, active: bool) -> None:
        self.conn.execute("""INSERT INTO topics(code,title,active) VALUES(?,?,?) ON CONFLICT(code) DO UPDATE SET title=excluded.title,active=excluded.active""", (code.strip(), title.strip(), int(active))); self.conn.commit()
    def set_setting(self, key: str, value: str) -> None:
        self.conn.execute("""INSERT INTO settings(key,value) VALUES(?,?) ON CONFLICT(key) DO UPDATE SET value=excluded.value""", (key.strip(), value)); self.conn.commit()
    def dashboard(self) -> dict[str, Any]:
        q = self.conn.execute
        sentiments = {r["sentiment"]: r["c"] for r in q("SELECT sentiment,COUNT(*) c FROM analysis_results GROUP BY sentiment").fetchall()}
        aspects: dict[str, int] = {}
        for r in q("SELECT aspects_json FROM analysis_results").fetchall():
            for a in json.loads(r["aspects_json"]): aspects[a] = aspects.get(a, 0) + 1
        return {"total_reviews":q("SELECT COUNT(*) c FROM reviews").fetchone()["c"],"analyzed_reviews":q("SELECT COUNT(*) c FROM analysis_results").fetchone()["c"],"negative_reviews":q("SELECT COUNT(*) c FROM analysis_results WHERE sentiment='negative'").fetchone()["c"],"critical_reviews":q("SELECT COUNT(*) c FROM analysis_results WHERE criticality='critical'").fetchone()["c"],"decision_count":q("SELECT COUNT(*) c FROM decisions").fetchone()["c"],"open_controls":q("SELECT COUNT(*) c FROM control_items WHERE status!='verified'").fetchone()["c"],"sentiments":sentiments,"aspects":dict(sorted(aspects.items(), key=lambda x:(-x[1],x[0])))}

ALIASES = {"отзыв":"text","текст":"text","review":"text","text":"text","дата":"date","date":"date","оценка":"rating","rating":"rating","автор":"author","author":"author","источник":"source","source":"source","канал":"channel","channel":"channel","id":"external_id","external_id":"external_id"}
def _header(value: Any) -> str:
    raw = str(value or "").strip().casefold(); return ALIASES.get(raw, raw)
def read_csv(path: str | Path) -> list[dict[str, Any]]:
    with Path(path).open("r", encoding="utf-8-sig", newline="") as f:
        sample=f.read(4096); f.seek(0)
        try: dialect=csv.Sniffer().sniff(sample,delimiters=";,\t,")
        except csv.Error: dialect=csv.excel; dialect.delimiter=";"
        return [{_header(k):v for k,v in row.items()} for row in csv.DictReader(f,dialect=dialect)]
def read_xlsx(path: str | Path) -> list[dict[str, Any]]:
    from openpyxl import load_workbook
    wb=load_workbook(path,read_only=True,data_only=True); ws=wb.active; rows=ws.iter_rows(values_only=True)
    try: headers=[_header(x) for x in next(rows)]
    except StopIteration: wb.close(); return []
    out=[]
    for values in rows:
        if any(v not in (None,"") for v in values): out.append({headers[i]:values[i] for i in range(min(len(headers),len(values)))})
    wb.close(); return out
def read_records(path: str | Path) -> list[dict[str, Any]]:
    suffix=Path(path).suffix.casefold()
    if suffix==".csv": return read_csv(path)
    if suffix in {".xlsx",".xlsm"}: return read_xlsx(path)
    raise ValueError(f"unsupported import format: {suffix}")
def build_report(stats: Mapping[str, Any]) -> str:
    lines=["# Аналитический отчёт FeedbackPro","",f"Сформирован: {datetime.utcnow().isoformat(timespec='seconds')} UTC","","## Сводка","",f"- Всего отзывов: **{stats['total_reviews']}**",f"- Проанализировано: **{stats['analyzed_reviews']}**",f"- Негативных: **{stats['negative_reviews']}**",f"- Критических: **{stats['critical_reviews']}**",f"- Решений: **{stats['decision_count']}**",f"- Открытых контрольных действий: **{stats['open_controls']}**","","## Тональность",""]
    for k in ("positive","neutral","negative"): lines.append(f"- {k}: {stats.get('sentiments',{}).get(k,0)}")
    lines += ["","## Темы/аспекты",""]
    for k,v in stats.get("aspects",{}).items(): lines.append(f"- {k}: {v}")
    lines += ["","## Апробация","","Показатели BEFORE/AFTER не задаются заранее и фиксируются только после фактической пилотной апробации.",""]
    return "\n".join(lines)

class FeedbackService:
    def __init__(self, db: Database): self.db=db; self.pipeline=AnalysisPipeline()
    def initialize(self) -> int: return self.db.migrate(2)
    def import_rows(self, rows: Iterable[Mapping[str, Any]], source_file: str | None = None, analyze: bool = True) -> dict[str,int]:
        stats={"inserted":0,"duplicates":0,"rejected":0,"analyzed":0}
        for row in rows:
            try: review=normalize_record(row,source_file=source_file); persisted,inserted=self.db.add_review(review)
            except (ValueError,TypeError): stats["rejected"]+=1; continue
            stats["inserted" if inserted else "duplicates"]+=1
            if analyze: self.db.save_analysis(self.pipeline.analyze(persisted)); stats["analyzed"]+=1
        return stats
    def import_file(self, path: str | Path, analyze: bool = True) -> dict[str,int]: return self.import_rows(read_records(path),Path(path).name,analyze)
    def analyze_all(self) -> int:
        reviews=self.db.list_reviews()
        for review in reviews: self.db.save_analysis(self.pipeline.analyze(review))
        return len(reviews)
    def create_decision(self,title:str,description:str,priority:str,review_id:int,due_date:date|None=None)->int: return self.db.create_decision(title,description,priority,[review_id],due_date)
    def set_control(self,decision_id:int,status:ControlStatus,due_date:date|None=None,outcome:str|None=None)->int: return self.db.set_control(decision_id,status,due_date,outcome)
    def dashboard(self)->dict[str,Any]: return self.db.dashboard()
    def report(self)->str: return build_report(self.dashboard())
