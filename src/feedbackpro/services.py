from __future__ import annotations

import csv
from datetime import date, datetime
from pathlib import Path
from typing import Any, Iterable, Mapping

from .analysis import AnalysisPipeline
from .domain import ControlStatus, normalize_record
from .storage import Database

ALIASES = {
    "отзыв": "text", "текст": "text", "review": "text", "text": "text",
    "дата": "date", "date": "date", "оценка": "rating", "rating": "rating",
    "автор": "author", "author": "author", "источник": "source", "source": "source",
    "канал": "channel", "channel": "channel", "id": "external_id", "external_id": "external_id",
}


def _header(value: Any) -> str:
    raw = str(value or "").strip().casefold()
    return ALIASES.get(raw, raw)


def read_csv(path: str | Path) -> list[dict[str, Any]]:
    with Path(path).open("r", encoding="utf-8-sig", newline="") as f:
        sample = f.read(4096)
        f.seek(0)
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=";,\t,")
        except csv.Error:
            dialect = csv.excel
            dialect.delimiter = ";"
        return [{_header(k): v for k, v in row.items()} for row in csv.DictReader(f, dialect=dialect)]


def read_xlsx(path: str | Path) -> list[dict[str, Any]]:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows = ws.iter_rows(values_only=True)
    try:
        headers = [_header(x) for x in next(rows)]
    except StopIteration:
        wb.close()
        return []
    out = []
    for values in rows:
        if any(v not in (None, "") for v in values):
            out.append({headers[i]: values[i] for i in range(min(len(headers), len(values)))})
    wb.close()
    return out


def read_records(path: str | Path) -> list[dict[str, Any]]:
    suffix = Path(path).suffix.casefold()
    if suffix == ".csv":
        return read_csv(path)
    if suffix in {".xlsx", ".xlsm"}:
        return read_xlsx(path)
    raise ValueError(f"unsupported import format: {suffix}")


def build_report(stats: Mapping[str, Any]) -> str:
    lines = [
        "# Аналитический отчёт FeedbackPro", "",
        f"Сформирован: {datetime.utcnow().isoformat(timespec='seconds')} UTC", "",
        "## Сводка", "",
        f"- Всего отзывов: **{stats['total_reviews']}**",
        f"- Проанализировано: **{stats['analyzed_reviews']}**",
        f"- Негативных: **{stats['negative_reviews']}**",
        f"- Критических: **{stats['critical_reviews']}**",
        f"- Решений: **{stats['decision_count']}**",
        f"- Открытых контрольных действий: **{stats['open_controls']}**",
        "", "## Тональность", "",
    ]
    for k in ("positive", "neutral", "negative"):
        lines.append(f"- {k}: {stats.get('sentiments', {}).get(k, 0)}")
    lines += ["", "## Темы/аспекты", ""]
    for k, v in stats.get("aspects", {}).items():
        lines.append(f"- {k}: {v}")
    lines += [
        "", "## Апробация", "",
        "Показатели BEFORE/AFTER не задаются заранее и фиксируются только после фактической пилотной апробации.", "",
    ]
    return "\n".join(lines)


class FeedbackService:
    def __init__(self, db: Database):
        self.db = db
        self.pipeline = AnalysisPipeline()

    def initialize(self) -> int:
        return self.db.migrate(2)

    def initialize_a3(self) -> int:
        return self.db.migrate(3)

    def import_rows(self, rows: Iterable[Mapping[str, Any]], source_file: str | None = None, analyze: bool = True) -> dict[str, int]:
        stats = {"inserted": 0, "duplicates": 0, "rejected": 0, "analyzed": 0}
        for row in rows:
            try:
                review = normalize_record(row, source_file=source_file)
                persisted, inserted = self.db.add_review(review)
            except (ValueError, TypeError):
                stats["rejected"] += 1
                continue
            stats["inserted" if inserted else "duplicates"] += 1
            if analyze:
                self.db.save_analysis(self.pipeline.analyze(persisted))
                stats["analyzed"] += 1
        return stats

    def import_file(self, path: str | Path, analyze: bool = True) -> dict[str, int]:
        return self.import_rows(read_records(path), Path(path).name, analyze)

    def analyze_all(self) -> int:
        reviews = self.db.list_reviews()
        for review in reviews:
            self.db.save_analysis(self.pipeline.analyze(review))
        return len(reviews)

    def create_decision(self, title: str, description: str, priority: str, review_id: int, due_date: date | None = None) -> int:
        return self.db.create_decision(title, description, priority, [review_id], due_date)

    def set_control(self, decision_id: int, status: ControlStatus, due_date: date | None = None, outcome: str | None = None) -> int:
        return self.db.set_control(decision_id, status, due_date, outcome)

    def dashboard(self) -> dict[str, Any]:
        return self.db.dashboard()

    def report(self) -> str:
        return build_report(self.dashboard())
